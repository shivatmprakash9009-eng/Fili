"""Prospecting helpers for contact hygiene, exports, and safe website fingerprints.

MX records confirm domain mail routing; they do not verify an individual mailbox.
Website scanning is restricted to public HTTP(S) destinations and bounded responses.
"""
from __future__ import annotations

import csv
import io
import ipaddress
import re
import socket
from typing import Any
from urllib.parse import urljoin, urlsplit, urlunsplit

PLACEHOLDER_EMAILS = {"", "unknown", "n/a", "none", "null"}
PLACEHOLDER_PHONES = ("555-01", "555-010", "xxxxx", "xxxx")
EXPORT_FIELDS = [
    "company_name", "contact_name", "contact_title", "contact_email", "phone",
    "email_verification_status", "website", "tech_stack", "location", "industry",
    "relevance_score", "platform", "source_type", "contact_verified",
    "email_subject", "email_body",
]


def clean_phone(value: Any) -> str:
    raw = str(value or "").strip()
    if not raw or any(token in raw.lower() for token in PLACEHOLDER_PHONES):
        return ""
    digits = re.sub(r"\D", "", raw)
    return raw if 7 <= len(digits) <= 15 else ""


def normalize_whatsapp_phone(value: Any) -> str:
    """Return E.164 digits only; never guess a country code for local numbers."""
    raw = str(value or "").strip()
    if not raw.startswith("+"):
        return ""
    digits = re.sub(r"\D", "", raw)
    if not 8 <= len(digits) <= 15 or digits.startswith("0"):
        return ""
    return digits


def email_domain(value: Any) -> str:
    email = str(value or "").strip().lower()
    if email in PLACEHOLDER_EMAILS or not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        return ""
    return email.rsplit("@", 1)[1].strip(".")


def verify_email_domain(value: Any, timeout: float = 2.5) -> dict[str, Any]:
    """Return domain MX status. This never claims that a mailbox is verified."""
    domain = email_domain(value)
    if not domain:
        return {"status": "invalid_format", "mx_found": False, "mailbox_verified": False}
    try:
        import dns.resolver  # type: ignore
        answers = dns.resolver.resolve(domain, "MX", lifetime=timeout)
        hosts = sorted({str(answer.exchange).rstrip(".") for answer in answers})
        return {
            "status": "mx_found" if hosts else "mx_not_found",
            "mx_found": bool(hosts),
            "mx_hosts": hosts,
            "mailbox_verified": False,
            "note": "MX confirms domain mail routing, not mailbox existence or deliverability.",
        }
    except ImportError:
        return {
            "status": "mx_lookup_unavailable", "mx_found": False, "mailbox_verified": False,
            "note": "Install dnspython to perform MX lookups.",
        }
    except Exception as exc:
        return {
            "status": "mx_not_confirmed", "mx_found": False, "mailbox_verified": False,
            "note": str(exc)[:180],
        }


def enrich_lead(lead: dict[str, Any], verify_mx: bool = True) -> dict[str, Any]:
    enriched = dict(lead)
    enriched["phone"] = clean_phone(enriched.get("phone"))
    email = str(enriched.get("contact_email") or "").strip()
    domain = email_domain(email)
    if domain:
        enriched["email_verification"] = (
            verify_email_domain(email) if verify_mx else
            {"status": "not_checked", "mx_found": False, "mailbox_verified": False}
        )
    else:
        enriched["contact_email"] = ""
        enriched["email_verification"] = {
            "status": "missing_or_placeholder", "mx_found": False, "mailbox_verified": False
        }
    enriched["email_domain"] = domain
    enriched["phone_available"] = bool(enriched["phone"])
    enriched["whatsapp_url"] = whatsapp_link(enriched)
    enriched["follow_up_sequence"] = build_follow_up_sequence(enriched)
    return enriched


def whatsapp_link(lead: dict[str, Any]) -> str:
    """Create a click-to-chat draft only when an explicit international phone is present."""
    from urllib.parse import quote
    number = normalize_whatsapp_phone(lead.get("phone"))
    if not number:
        return ""
    contact = str(lead.get("contact_name") or "there").strip()
    company = str(lead.get("company_name") or "your team").strip()
    hook = str(lead.get("key_details") or "").strip()
    message = f"Hi {contact}, reaching out regarding {company}."
    if hook:
        message += f" {hook}"
    message += " Would a brief conversation be useful?"
    return f"https://wa.me/{number}?text={quote(message)}"


def validate_public_http_url(value: Any) -> str:
    """Validate a public HTTP(S) URL and reject local/private/reserved destinations."""
    raw = str(value or "").strip()
    if not raw or len(raw) > 2048:
        raise ValueError("Enter a valid website URL (maximum 2048 characters).")
    candidate = raw if "://" in raw else f"https://{raw}"
    try:
        parsed = urlsplit(candidate)
        if parsed.scheme.lower() not in {"http", "https"}:
            raise ValueError("Only HTTP and HTTPS websites can be scanned.")
        if parsed.username or parsed.password or not parsed.hostname:
            raise ValueError("Website URL must not contain credentials and must include a hostname.")
        port = parsed.port
        host = parsed.hostname.rstrip(".").lower()
        if host in {"localhost"} or host.endswith((".localhost", ".local", ".internal", ".test")):
            raise ValueError("Local and internal hosts cannot be scanned.")
        try:
            addresses = [ipaddress.ip_address(host)]
        except ValueError:
            ascii_host = host.encode("idna").decode("ascii")
            resolved = socket.getaddrinfo(ascii_host, port or (443 if parsed.scheme == "https" else 80),
                                          type=socket.SOCK_STREAM)
            addresses = [ipaddress.ip_address(item[4][0]) for item in resolved]
        if not addresses or any(not address.is_global for address in addresses):
            raise ValueError("Only publicly routable website addresses can be scanned.")
        netloc = parsed.netloc
        safe = urlunsplit((parsed.scheme.lower(), netloc, parsed.path or "/", parsed.query, ""))
        return safe
    except ValueError:
        raise
    except (OSError, UnicodeError) as exc:
        raise ValueError("Website hostname could not be safely resolved.") from exc


def tech_stack_signals(url: str, html: str = "", headers: dict[str, str] | None = None) -> list[str]:
    """Best-effort public HTML/header fingerprint; signals are not guaranteed."""
    parsed = urlsplit(url if "://" in url else f"https://{url}")
    domain = (parsed.hostname or "").lower()
    text = (html or "").lower()
    header_text = " ".join(f"{k}:{v}" for k, v in (headers or {}).items()).lower()
    haystack = text + " " + header_text
    checks = {
        "Shopify": ("cdn.shopify.com", "shopify.theme", "x-shopid"),
        "WordPress": ("wp-content", "wp-includes", "wordpress"),
        "WooCommerce": ("woocommerce",),
        "Webflow": ("webflow.js", "webflow.com"),
        "Wix": ("wixstatic.com", "wix.com"),
        "Next.js": ("_next/static", "__next_data__"),
        "React": ("react-dom", "data-reactroot", "__react"),
        "HubSpot": ("js.hs-scripts.com", "hubspot"),
        "Google Tag Manager": ("googletagmanager.com",),
        "Cloudflare": ("cf-ray", "cloudflare"),
        "Stripe": ("js.stripe.com", "stripe.com/v3"),
    }
    found = [name for name, needles in checks.items() if any(n in haystack for n in needles)]
    if domain.endswith("webflow.io"):
        found.append("Webflow")
    return sorted(set(found))


def fetch_tech_stack(url: str, timeout: float = 5.0, max_response_bytes: int = 1_000_000) -> dict[str, Any]:
    """Fetch a bounded public page, revalidating every redirect target against SSRF."""
    try:
        import httpx
        target = str(url or "").strip()
        headers: dict[str, str] = {}
        status_code = 0
        body = b""
        final_url = ""
        for redirect_count in range(4):
            target = validate_public_http_url(target)
            with httpx.Client(timeout=timeout, follow_redirects=False, trust_env=False,
                              headers={"User-Agent": "FiliTechSignals/1.0"}) as client:
                with client.stream("GET", target) as response:
                    status_code = response.status_code
                    if status_code in {301, 302, 303, 307, 308}:
                        location = response.headers.get("location")
                        if not location:
                            raise ValueError("Website returned a redirect without a destination.")
                        if redirect_count == 3:
                            raise ValueError("Website exceeded the redirect limit.")
                        target = urljoin(target, location)
                        continue
                    length = response.headers.get("content-length")
                    if length and int(length) > max_response_bytes:
                        raise ValueError("Website response is larger than the 1 MB scan limit.")
                    chunks = []
                    total = 0
                    for chunk in response.iter_bytes():
                        total += len(chunk)
                        if total > max_response_bytes:
                            raise ValueError("Website response is larger than the 1 MB scan limit.")
                        chunks.append(chunk)
                    body = b"".join(chunks)
                    headers = dict(response.headers)
                    final_url = str(response.url)
                    break
        text = body.decode("utf-8", errors="replace")
        signals = tech_stack_signals(final_url, text, headers)
        return {
            "status": "scanned", "signals": signals, "http_status": status_code,
            "final_url": final_url,
            "note": "Best-effort public fingerprint only; this is not a complete infrastructure inventory.",
        }
    except Exception as exc:
        return {"status": "unavailable", "signals": [], "note": str(exc)[:180]}


def build_follow_up_sequence(lead: dict[str, Any]) -> list[dict[str, Any]]:
    """Create three editable drafts. These are never sent automatically."""
    company = str(lead.get("company_name") or "your team")
    contact = str(lead.get("contact_name") or "there")
    subject = str(lead.get("email_subject") or f"Idea for {company}")
    first = str(lead.get("email_body") or f"Hi {contact},\n\nI wanted to share a relevant idea for {company}.")
    return [
        {"step": 1, "day": 1, "label": "Value hook", "subject": subject, "body": first},
        {"step": 2, "day": 3, "label": "Insight follow-up",
         "subject": f"One more idea for {company}",
         "body": f"Hi {contact},\n\nFollowing up with one additional thought for {company}. A short example or relevant result may help assess whether this is worth exploring.\n\nWould a brief overview be useful?"},
        {"step": 3, "day": 7, "label": "Polite close",
         "subject": f"Should I close the loop, {contact}?",
         "body": f"Hi {contact},\n\nI don't want to clutter your inbox. Should I close the loop for now, or would a short conversation about {company}'s priorities be useful?\n\nIf this isn't relevant, reply 'no' and I won't follow up."},
    ]


def _spreadsheet_safe(value: Any) -> str | int | float:
    """Neutralize spreadsheet formula prefixes in untrusted lead values."""
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = str(value)
    if text.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _export_rows(leads: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for lead in leads:
        row = dict(lead)
        row["email_verification_status"] = (row.get("email_verification") or {}).get("status", "not_checked")
        row["tech_stack"] = ", ".join(row.get("tech_stack") or row.get("technologies_used") or [])
        row["contact_verified"] = bool(row.get("contact_verified", False))
        rows.append({field: _spreadsheet_safe(row.get(field, "")) for field in EXPORT_FIELDS})
    return rows


def csv_export(leads: list[dict[str, Any]]) -> str:
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=EXPORT_FIELDS, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(_export_rows(leads))
    return output.getvalue()


def excel_export(leads: list[dict[str, Any]]) -> bytes:
    """Build an XLSX CRM export without interpreting untrusted lead cells as formulas."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fili CRM"
    sheet.append(EXPORT_FIELDS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="202020")
    for row in _export_rows(leads):
        sheet.append([row.get(field, "") for field in EXPORT_FIELDS])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        max_len = max((len(str(cell.value or "")) for cell in column_cells), default=10)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 12), 48)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
